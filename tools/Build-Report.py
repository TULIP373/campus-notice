"""Rebuild the five-experiment report from the retained course template.

Run with the Codex bundled Python. No credentials are read. --pages accepts
the actual bookmark/page map collected by Word; template geometry and opaque
package parts remain unchanged. Visual QA is a separate Word/PDF pass.
"""
from pathlib import Path
from copy import deepcopy
import argparse
import hashlib
import json
import re
import textwrap
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / 'docs/report-template.docx'
if not REFERENCE.exists():
    REFERENCE = ROOT / '.runtime/report-reference/template.docx'
OUTPUT = ROOT / 'deliverables/202232110113余家祥.docx'
EVIDENCE = ROOT / 'reports/2026-10-08'
QA = EVIDENCE / 'documentQA'
JAVA = 'src/main/java/com/example/campusnotice/'
TITLES = ['实验一 后端分层与用户管理', '实验二 前后端分离与课程通知',
          '实验三 Git协作与构建检查', '实验四 登录认证与角色权限',
          '实验五 课程资料检索与模型问答']
SECTIONS = ['项目介绍', '相关开发技术', '项目设计（概要设计、详细设计、数据库设计）',
            '关键代码', '运行结果', '心得（使用 AI 情况）']


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_json(name):
    path = EVIDENCE / name
    return json.loads(path.read_text(encoding='utf-8-sig')) if path.exists() else {}


def excerpt(path, start, end=None, count=14):
    lines = (ROOT / path).read_text(encoding='utf-8-sig').splitlines()
    idx = next(i for i, line in enumerate(lines) if start in line)
    if end:
        stop = next(i for i in range(idx + 1, len(lines)) if end in lines[i])
        chosen = lines[idx:stop]
    else:
        chosen = lines[idx:idx + count]
    return '\n'.join(chosen)


class Builder:
    def __init__(self):
        self.doc = Document(REFERENCE)
        self.source_paragraphs = [deepcopy(p._p) for p in self.doc.paragraphs]
        self.section_proto = self.source_paragraphs[13]
        self.title_proto = self.source_paragraphs[12]
        # Keep the source cover and its explicit start-of-contents page.
        for child in list(self.doc._element.body):
            if child.tag != qn('w:sectPr'):
                self.doc._element.body.remove(child)
        for para in self.source_paragraphs[:11]:
            self.doc._element.body.insert(len(self.doc._element.body)-1, para)
        if 'Title' not in self.doc.styles:
            from docx.enum.style import WD_STYLE_TYPE
            title_style = self.doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
            title_style.base_style = self.doc.styles['Normal']
        self.doc.paragraphs[4].style = self.doc.styles['Title']
        for index, text in [(6, '学号：202232110113'), (7, '姓名：余家祥'), (8, '班级：软外241')]:
            p = self.doc.paragraphs[index]
            old = deepcopy(p.runs[0]._r.rPr) if p.runs else None
            p.clear()
            r = p.add_run(text)
            if old is not None:
                r._r.insert(0, old)
        self.doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
        self.doc.styles['Title'].paragraph_format.space_before = Pt(0)
        self.doc.styles['Title'].paragraph_format.space_after = Pt(0)
        self.font(self.doc.styles['Normal'].font, 10.5)
        self.bookmark = 100
        self.num_ids = []
        numbering = self.doc.part.numbering_part.element
        current = [int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]
        for i in range(5):
            num_id = max(current + [0]) + 1 + i
            num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
            abstract = OxmlElement('w:abstractNumId'); abstract.set(qn('w:val'), '0'); num.append(abstract)
            override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
            start = OxmlElement('w:startOverride'); start.set(qn('w:val'), '1'); override.append(start); num.append(override)
            numbering.append(num); self.num_ids.append(num_id)

    def font(self, font, size=10.5, code=False):
        font.name = 'Consolas' if code else '宋体'; font.size = Pt(size)
        font.color.rgb = RGBColor(0, 0, 0); font.strike = False
        if hasattr(font, '_element'):
            pr = font._element
            if pr.tag == qn('w:rPr'):
                fonts = pr.find(qn('w:rFonts'))
                if fonts is None: fonts = OxmlElement('w:rFonts'); pr.insert(0, fonts)
                fonts.set(qn('w:eastAsia'), '宋体')

    def para(self, text, bold=False, code=False):
        p = self.doc.add_paragraph(style='Normal')
        p.paragraph_format.space_after = Pt(4 if not code else 0)
        p.paragraph_format.line_spacing = 1.12 if not code else 1.0
        p.paragraph_format.widow_control = True
        if code:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(text); self.font(r.font, 9 if code else 10.5, code); r.bold = bold
        return p

    def heading(self, text, num_id=None, bookmark=None, page=False):
        from docx.text.paragraph import Paragraph
        elm = deepcopy(self.section_proto if num_id else self.title_proto)
        self.doc._element.body.insert(len(self.doc._element.body)-1, elm)
        p = Paragraph(elm, self.doc._body); p.clear()
        r = p.add_run(text); self.font(r.font, 18); r.bold = False
        p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(6); p.paragraph_format.page_break_before = page
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if num_id else WD_ALIGN_PARAGRAPH.CENTER
        if num_id:
            np = elm.get_or_add_pPr().get_or_add_numPr()
            np.get_or_add_numId().val = num_id; np.get_or_add_ilvl().val = 0
        else:
            np = elm.get_or_add_pPr().find(qn('w:numPr'))
            if np is not None: elm.get_or_add_pPr().remove(np)
        if bookmark:
            a = OxmlElement('w:bookmarkStart'); a.set(qn('w:id'), str(self.bookmark)); a.set(qn('w:name'), bookmark)
            b = OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'), str(self.bookmark)); self.bookmark += 1
            elm.insert(1, a); elm.append(b)
        return p

    def code(self, path, start, end=None, count=14):
        label = self.para('源码位置：' + path)
        label.alignment = WD_ALIGN_PARAGRAPH.LEFT
        label.paragraph_format.keep_with_next = True
        raw = textwrap.dedent(excerpt(path, start, end, count))
        displayed = []
        for line in raw.splitlines():
            # Keep source lines intact; Word soft-wraps visually without inserting
            # paragraph breaks inside identifiers or string literals.
            displayed.append(self.para(line, code=True))
        for p in displayed[:-1]: p.paragraph_format.keep_with_next = True

    def table(self, headers, rows, widths):
        t = self.doc.add_table(rows=1, cols=len(headers)); t.autofit = False
        for i, (h, w) in enumerate(zip(headers, widths)):
            t.columns[i].width = Cm(w); t.rows[0].cells[i].text = h
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row): cell.text = str(value)
        pr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e = OxmlElement('w:'+edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
        pr.append(borders)
        for index, row in enumerate(t.rows):
            no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.keep_with_next = index == 0
                    for r in p.runs: self.font(r.font, 9.5); r.bold = index == 0
                if index == 0:
                    s = OxmlElement('w:shd'); s.set(qn('w:fill'), 'F2F2F2'); cell._tc.get_or_add_tcPr().append(s)
            if index == 0:
                repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        self.para('')

    def image(self, name, caption, crop=None):
        path = EVIDENCE / name
        if not path.exists():
            self.para('截图尚未形成：' + name); return
        p = self.doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        from PIL import Image
        with Image.open(path) as picture:
            width, height = picture.size
        left, top, right, bottom = crop or (0, 0, width, height)
        visible_width, visible_height = right-left, bottom-top
        display_width = min(14.4, 12.8 * visible_width / visible_height)
        picture = p.add_run().add_picture(str(path), width=Cm(display_width),
                                         height=Cm(display_width*visible_height/visible_width))
        if crop:
            # Word crops the embedded original; no pixels are altered or redrawn.
            fill = picture._inline.find('.//' + qn('pic:blipFill'))
            rect = OxmlElement('a:srcRect')
            for side, value in [('l',left/width), ('t',top/height),
                                ('r',(width-right)/width), ('b',(height-bottom)/height)]:
                rect.set(side, str(round(value * 100000)))
            fill.insert(1, rect)
            picture._inline.docPr.set('descr', f'真实截图 {name}；Word裁剪局部 {crop}；原图完整嵌入')
        cap = self.para(caption); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs: self.font(r.font, 10)

    def footer(self):
        p = self.doc.sections[0].footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(); self.font(r.font, 10)
        begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin'); r._r.append(begin)
        instr = OxmlElement('w:instrText'); instr.text = ' PAGE '; r._r.append(instr)
        separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate'); r._r.append(separate)
        r = p.add_run('1'); self.font(r.font, 10)
        end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); r._r.append(end)
        update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); self.doc.settings.element.append(update)

    def save(self):
        OUTPUT.parent.mkdir(parents=True, exist_ok=True); QA.mkdir(parents=True, exist_ok=True)
        self.doc.save(OUTPUT)
        with zipfile.ZipFile(REFERENCE) as z: source = {n:z.read(n) for n in z.namelist()}
        with zipfile.ZipFile(OUTPUT) as z: final = {n:z.read(n) for n in z.namelist()}
        preserve = [n for n in source if n.startswith(('customXml/', 'docProps/', 'word/theme/', 'word/fontTable.xml')) or n == '_rels/.rels']
        for name in preserve: final[name] = source[name]
        with zipfile.ZipFile(OUTPUT, 'w', zipfile.ZIP_DEFLATED) as z:
            for name, value in final.items(): z.writestr(name, value)
        from lxml import etree
        ref_xml = etree.fromstring(source['word/document.xml']); out_xml = etree.fromstring(final['word/document.xml'])
        def geo(root):
            section = root.find('.//'+qn('w:sectPr'))
            return {c.tag: dict(c.attrib) for c in section if c.tag in (qn('w:pgSz'),qn('w:pgMar'))}
        facts = {'template_source':str(REFERENCE.relative_to(ROOT)), 'template_sha256':digest(REFERENCE), 'report_sha256':digest(OUTPUT),
                 'template_unchanged':digest(REFERENCE)=='8b88124b5d66f0039b9d19deb07a22ca463b54fd284834de3a7fcafd18dd5483',
                 'preserve_only_parts':preserve, 'preserve_only_equal':all(source[n]==final[n] for n in preserve),
                 'geometry_equal':geo(ref_xml)==geo(out_xml), 'experiments':5, 'sections_per_experiment':6,
                 'added_parts':[n for n in final if n not in source], 'render_status':'PENDING',
                 'intentional_changes':['身份填写','五实验正文及真实截图','目录真实页码','PAGE页脚与打开时更新字段']}
        (QA/'structure-check.json').write_text(json.dumps(facts,ensure_ascii=False,indent=2),encoding='utf-8')
        print(OUTPUT)


def make(pages):
    b = Builder()
    b.para('课程通知平台  五个实验合订报告', bold=True)
    b.para('本报告记录从用户管理后端、Vue课程通知界面、Git协作到权限控制和课程资料问答的实现过程。运行结果对应本地新项目的真实日志与界面。演示课程资料用于功能验证，不代表教师发布的实际任务。')
    for i, title in enumerate(TITLES,1): b.para(title + '  ' + str(pages.get('exp'+str(i),'待核对')))
    b.para('每个实验依次包含项目介绍、相关开发技术、项目设计、关键代码、运行结果、心得六栏目。')
    for i, title in enumerate(TITLES,1):
        # Body experiments flow naturally; only the source cover/contents/body
        # boundary forces a new page. This avoids isolated reflection pages.
        b.heading(title,bookmark='exp'+str(i),page=i == 1)
        for section_index, section in enumerate(SECTIONS):
            b.heading(section, b.num_ids[i-1])
            content(b,i,section_index)
            if section_index == 5:
                records = {
                    1: ('Maven环境、分层代码及用户测试', 'Java17、SpringBoot3、真实MySQL、逐层实现并验证CRUD', '构建、真实接口、重启持久化与密码保护断言'),
                    2: ('Vue组件、Axios、接口文档与联调', '三课程通知和资料、真实API、明确输入及错误边界', '前后端构建、CORS、真实浏览器及数据库增改删'),
                    3: ('分支冲突方案、PR模板和CI检查', '每个实际步骤中文提交，保留真实冲突，不伪造远程结果', '双父合并、冲突日志、提交diff和本地CI命令'),
                    4: ('登录、安全链、JWT与角色界面', 'BCrypt、教师写权限、内存令牌、秘密不提交', '登录及401/403、篡改、有效签名的过期令牌和刷新重登'),
                    5: ('资料检索、提示、SDK桥及问答原型', '只答本课程原文、未知不猜、引用可核对、故障不伪造', '真实SDK、三次提示练习、课程隔离和注入、模型停止错误及原文比对'),
                }
                scope, prompt, checks = records[i]
                b.para('AI使用记录：Codex辅助'+scope+'；提示要求“'+prompt+'”；产出为相应代码、配置和检查建议，通过'+checks+'核验。本人仍须阅读、理解并复述关键流程，自动验证不等于本人已逐行理解。')
    b.footer(); b.save()


def content(b, experiment, section):
    if experiment == 1:
        if section == 0:
            b.para('实验一建立课程通知平台后端，先实现用户增删改查，为后续课程数据和身份认证提供基础。系统采用独立MySQL保存数据，控制器接收JSON，服务层检查业务条件，Mapper执行参数化SQL。')
            b.para('实施顺序为保留IDEA创建的Spring Boot项目、固定JDK17与Maven Wrapper、配置数据源、编写用户表和实体、增加Mapper、Service与Controller，再验证接口和重启后的数据。')
        elif section == 1:
            b.para('后端采用Java17、Spring Boot3、Spring Web、MyBatis与MySQL；Maven Wrapper用于一致构建。BCrypt负责密码单向哈希，JUnit与Mockito验证业务边界。实际数据库运行于本机3307端口，后端运行于8080端口。')
            b.para('MyBatis的#{参数}绑定让用户名等内容作为数据进入SQL。异常处理将重复用户名、非法输入、不存在记录分别转换为409、400、404；用户响应不暴露passwordHash。')
        elif section == 2:
            b.para('概要设计：请求进入UserController，调用UserService，再由UserMapper访问users表。Entity描述持久化结构，UserRequest描述输入，ApiExceptionHandler统一错误消息。')
            b.para('详细设计：新增用户先校验用户名、姓名、角色与密码，计算哈希后插入并返回查询结果；更新资料时未填写新密码则保留原哈希；删除不存在用户返回404。正整数ID约束在服务层校验。')
            b.para('数据库设计：users使用自增主键，username有唯一索引；中文姓名采用utf8mb4。密码列只保存哈希，created_at由数据库默认生成。')
            b.table(['字段','类型','约束与用途'], [('id','BIGINT','自增主键'),('username','VARCHAR(50)','唯一，不为空'),('name','VARCHAR(100)','中文姓名'),('password_hash','VARCHAR(100)','BCrypt哈希'),('role','VARCHAR(20)','STUDENT或TEACHER'),('created_at','TIMESTAMP','创建时间')],[3.2,3.3,8])
        elif section == 3:
            b.code(JAVA+'service/UserService.java','public User create','public User update')
            b.code(JAVA+'mapper/UserMapper.java','@Select("SELECT id, username, name, password_hash, role, created_at FROM users WHERE id',count=2)
            b.para('服务层把数据库唯一约束冲突映射为409，避免把原始SQL异常直接返回给界面。Mapper的id绑定和密码哈希逻辑分别承担查询安全与密码存储安全。')
        elif section == 4:
            b.para('真实接口验证完成新增、读取、中文更新和删除。重复用户名返回409，非法输入返回400，删除后再次查询返回404。数据库重启流程另有状态和端口证据，后续读取持久化记录通过；这不是内存存储验证。')
            b.table(['检查','实际观察','证据'], [('CRUD与中文','PASS','api-check.json'),('重复及非法输入','409及400','e1-crud-api.json'),('重启持久化','通过','persistence-state.json、e1-restart.json'),('密码保护测试','故意污染后断言失败，恢复后通过','e1-password-guard-failure.log、e1-unit-pass.log')],[3.2,4,7.3])
            b.para('阶段一构建通过；截至实验五的累计Java检查为40条、失败0、错误0。日志保存在reports/2026-10-08/e5-ask-build.log。不同阶段测试数量不同，不能把最终40条全部归为实验一。')
        else:
            b.para('最初测试对数据库约束行为的假设有误，独立测试作者依据验收标准修正并保留失败记录。“构建成功”与“业务正确”需要不同证据。本人需解释请求如何进入数据库、重复用户名为何返回409，以及JSON为何不能含密码哈希。')
    elif experiment == 2:
        if section == 0:
            b.para('实验二将后端扩展为课程通知平台，提供按课程和关键词查找通知、查看资料和提交要求的界面。教师可以新增、修改和删除通知；学生可阅读标题、正文、报告格式、截止时间与来源。演示数据包含三门课程，每门都有通知和学习资料。')
            b.para('步骤为建立课程和通知表、实现REST JSON接口、配置CORS与接口文档、创建Vue组件和Axios客户端，再完成学生详情和教师表单，分别构建后进行真实浏览器联调。')
        elif section == 1:
            b.para('Vue3负责组件状态和页面渲染，Vite负责开发和生产构建，Axios发送HTTP请求。Spring的@RestController自动将对象序列化为JSON，Jakarta Validation检查通知输入。')
            b.para('前后端分别运行在5173和8080端口。CORS只允许规定的本地前端来源；它控制浏览器跨域访问，接口业务权限在实验四由Spring Security补充。接口约定写入OpenAPI文档。')
        elif section == 2:
            b.para('概要设计：App组合CourseList、NoticeList、NoticeDetail与NoticeForm。切换课程触发接口查询，列表条目与详情共享选择状态。表单保存成功后重新拉取数据，避免页面仅修改本地显示而数据库未更新。')
            b.para('详细设计：GET /api/courses获取课程；GET /api/notices支持courseId和q；GET /api/notices/{id}读取详情；POST、PUT和DELETE完成教师维护。空截止时间显示待教师通知，错误提示与空列表状态分别处理。')
            b.para('数据库设计：courses与notices是一对多；notices.course_id有外键和索引。通知类型限制为NOTICE或MATERIAL，提交要求和报告格式独立保存，便于详情展示和后续资料检索。')
            b.table(['表或字段','关系及用途'], [('courses','课程名称唯一，保存介绍'),('notices.course_id','关联courses.id'),('title / content','通知标题与正文'),('submission_requirements / report_format','提交要求与报告格式'),('deadline / source / type','截止时间、来源、通知或资料类型')],[5.5,9])
        elif section == 3:
            b.code(JAVA+'controller/NoticeController.java','@GetMapping','@GetMapping("/{id}")')
            b.code('frontend/src/api.js','export const api = axios.create',count=5)
            b.para('上述Controller摘录包含实验四加入的权限注解，反映最终源码。请求参数按courseId和q转交服务层，Axios基础URL支持环境变量配置。')
        elif section == 4:
            b.para('真实浏览器联调完成三门课切换、通知和资料选择、详情查看，以及教师表单新增和修改。修改提交要求后刷新并重新选课，读取到数据库中的更新值；实际DELETE另由接口测试通过。证据e2-browser-check.json记录了界面验证范围。')
            b.para('CORS允许http://localhost:5173，其他来源返回403且无允许头。frontend-build.log与后端日志记录分别构建成功。')
            b.image('e2-student.png','图2 1 实验二学生课程通知详情 真实本地界面')
            b.image('e2-teacher.png','图2 2 实验二教师编辑通知的正文与提交要求（真实截图局部）', (304,350,1262,884))
            b.image('e2-teacher.png','图2 3 实验二报告格式、截止时间、来源和保存按钮（同一真实截图局部）', (304,889,1262,1234))
        else:
            b.para('浏览器发现的问题需要结合真实请求和数据库重查，截图不能代替接口测试。个人学习重点是组件如何通知父组件重查、刷新后为何保留修改，以及CORS错误和业务错误的区别；演示资料必须明确标识。')
    elif experiment == 3:
        if section == 0:
            b.para('实验三使用本地Git保存开发过程，完成feature分支、真实冲突与合并，增加PR模板及构建检查。远程推送待本人完成，当前无GitHub仓库链接或在线PR。')
            b.para('先创建feature/commit-convention并与主分支分别修改同一行，保留冲突日志，解决后提交dec915d，再加入PR模板与CI，执行本地构建命令。')
        elif section == 1:
            b.para('Git分支隔离变更，merge整合历史。GitHub Actions配置Java17 Maven verify和Vue构建两个任务；PR模板说明变更、验证与风险。')
            b.para('CI文件存在只说明自动化流程已配置。未推送时无法证明云端工作流通过，因此本报告区分本地命令结果和云端执行状态。')
        elif section == 2:
            b.para('概要设计：实际步骤形成带实验编号的中文提交；两分支修改开发约定的同一行产生CONFLICT，随后保留两项规则再合并。')
            b.para('详细设计：提交摘要既包含实验编号和中文动作，又要求验证相关改动并保存证据。PR模板要求描述验证方式，CI使用只读contents权限并关闭checkout凭据持久化。')
            b.para('数据库设计：本实验不新增业务数据库表。版本库保存源码与SQL脚本，运行中的数据库文件、密钥、依赖缓存不作为版本数据提交。')
            b.table(['材料','实际状态'], [('feature/commit-convention','本地真实分支'),('dec915d','实际冲突解决与合并提交'),('.github/PULL_REQUEST_TEMPLATE.md','已编写'),('.github/workflows/ci.yml','已配置前后端两个任务'),('远程PR与云端CI','未推送，未执行')],[6.2,8.3])
        elif section == 3:
            b.code('.github/workflows/ci.yml','      - uses: actions/setup-java@v6',count=6)
            b.code('.github/workflows/ci.yml','      - run: npm ci',count=2)
            b.para('真实冲突来自docs/开发约定.md：主分支要求验证后留证，feature分支要求摘要包含实验编号。冲突标记保存在git-conflict.log；最终约定吸收两条规则，而不是删去一方要求。')
        elif section == 4:
            history = (EVIDENCE/'git-history.txt').read_text(encoding='utf-8-sig')
            b.para(f'git-history.txt的阶段记录包含{len(history.splitlines())}条提交，全部来自真实操作；后续安全与问答实现继续形成新提交。真实冲突日志显示CONFLICT (content)，合并提交为dec915d。')
            b.para('本地后端阶段测试14条通过，前端npm ci和npm run build最终通过。首次前端重装依赖被运行中的Vite占用原生模块阻止，停止进程后重跑成功；失败保存在e3-vite-lock-failure.log，最终通过以e3-local-ci-frontend.log为准。')
            b.table(['验证','结果','证据'], [('分支冲突','真实发生并解决','git-conflict.log、git-check.json'),('本地后端检查','通过','e3-local-ci-backend.log'),('本地前端检查','最终通过','e3-local-ci-frontend.log'),('远程操作','未执行','git-check.json')],[3,3.5,8])
        else:
            b.para('冲突日志、提交图和构建日志用于核对实际Git操作；配置文件不能证明云端CI通过。依赖锁定失败说明应等命令结束再记录结果。本人后续需阅读diff和测试，再推送自己的仓库并补真实URL。')
    elif experiment == 4:
        if section == 0:
            b.para('实验四为课程通知平台增加登录和角色权限。教师维护用户和通知，学生阅读课程与资料并使用问答。密码保存BCrypt哈希，JWT带有效期，前端令牌只在当前页面内存保存。')
            b.para('步骤为创建角色权限表、初始化本地演示账户、实现密码登录、签发与验证JWT、加入过滤器和@PreAuthorize、调整前端登录状态，再验证401、403、篡改、过期与学生写入边界。')
        elif section == 1:
            b.para('Spring Security使用无会话Bearer认证，Nimbus JWT负责HS256签发和校验，BCrypt进行密码比对，@EnableMethodSecurity启用方法权限。权限由数据库读取后进入令牌声明。')
            b.para('签名密钥只从私有配置或环境变量读取；演示账户初始化脚本生成密码哈希。Vue通过Axios拦截器发送Authorization，刷新页面后内存状态清空，需要重新登录。')
        elif section == 2:
            b.para('概要设计：登录接口核对账号密码并返回令牌和角色；安全过滤器验证签名、签发者与有效期；角色和权限转换为GrantedAuthority。接口过滤器先拒绝无权限写入，控制器注解再进行方法授权。')
            b.para('详细设计：登录失败、无令牌、过期和篡改返回401；已登录学生执行教师操作返回403。对写入先检查权限，避免无权限请求先触发MVC字段校验而返回400。前端authEpoch用于避免旧请求401误清除新会话。')
            b.para('数据库设计：users.role关联roles.code；roles和permissions通过role_permissions关联。一名用户一种角色，一个角色可拥有多个权限。')
            b.table(['角色','读取权限','维护权限','问答权限'], [('STUDENT','course:read / notice:read','无','ask:use'),('TEACHER','course:read / notice:read','notice:write / user:write','ask:use')],[2.5,4.6,4.6,2.8])
        elif section == 3:
            b.code(JAVA+'config/SecurityConfig.java','public JwtDecoder jwtDecoder',end='    @Bean')
            b.code(JAVA+'controller/NoticeController.java','    @PostMapping','    @PutMapping')
            b.code('frontend/src/api.js',"let token = ''",count=10)
            b.para('JWT解码器同时校验签名与campus-notice签发者；方法注解明确教师和notice:write双重要求。前端不把令牌写入localStorage或sessionStorage。')
        elif section == 4:
            b.para('security-check.json记录教师和学生登录成功，错误密码与匿名请求返回401，学生管理与通知增改删返回403，篡改JWT返回401，签名有效但已过期JWT返回401。SQL注入式输入被当作字面数据；问答输入长度0和1001被拒绝，1和1000边界通过。')
            b.para('本阶段Java安全及业务检查累计23条通过。后续实验五累计40条通过。e4-permission-order-failure.json保留权限过滤顺序修正前的真实失败，修正后重跑。截图记录登录后的实际角色界面，刷新清会话已在浏览器核对。')
            b.image('e4-student.png','图4 1 学生登录后查看课程通知（真实截图局部）', (0,0,1316,500))
            b.image('e4-teacher.png','图4 2 登录后的教师身份和维护入口（真实截图局部）', (0,0,1316,348))
            b.image('e4-teacher.png','图4 3 教师新增通知表单及保存入口（同一真实截图局部）', (304,351,1262,1234))
        else:
            b.para('真实问题是学生写入先被字段校验拦截，返回400而非403；教师写权限前置到过滤器后恢复预期。权限须由服务端拒绝，隐藏按钮不能形成安全边界。本人需解释BCrypt与JWT签名的区别，以及过期401和权限不足403。')
    else:
        if section == 0:
            b.para('实验五增加课程资料问答：按当前课程检索通知与资料原文，再将System规则和User问题、片段传给真实模型，回答附原文编号引用。检索、校验与模型接口已经实现；报告对真实模型调用结果按当前证据记录。')
            b.para('步骤为设计System/User提示、实现SDK访问桥、资料切块和稀疏向量检索、加入问答接口和Vue组件，再验证课程隔离、无资料拒答、注入、引用和模型故障。GPT Plus登录接入与独立API Key的计费路径分别说明。')
        elif section == 1:
            b.para('RAG采用中文字符和相邻双字TF-IDF稀疏向量，余弦相似度检索当前课程前三块，不调用额外embedding服务。文本切块450字、重叠80字，低于0.12相似度的片段过滤。')
            b.para('Java ModelClient通过HTTP连接本地模型桥，由桥选择官方Codex SDK或本人显式配置的兼容API模式。实际验证使用官方SDK/CLI的ChatGPT登录；兼容API模式尚未实际验证。模型不可用时返回明确错误，不能用固定答案代替生成结果。')
        elif section == 2:
            b.para('概要设计：AskController接收courseId和question，AskService验证输入和课程范围，RagService从当前课程通知构造原文片段，ModelClient发送提示，再校验引用并返回回答与来源。')
            b.para('详细设计：先隔离课程再向量化；问题1至1000字，命中泄密或读取文件指令直接拒绝；检索无结果则说明无法可靠回答。System要求问题和资料均作为不可信数据，回答每项事实附[1]式引用。引用为空或编号超出来源范围返回502。编号检查不等于逐句事实准确性证明，仍需对照原文。')
            b.para('数据库设计：复用courses和notices，当前未新建向量库。通知正文、提交要求、格式、截止时间和来源在查询时组成文本。中文稀疏向量只在本次请求中构造；API返回noticeId、课程、标题、来源、片段与得分，便于观察检索过程。')
            b.table(['环节','规则'], [('切块','450字，80字重叠'),('课程隔离','先过滤courseId再建向量'),('向量','字与双字TF-IDF，双字权重2'),('检索','余弦分数至少0.12，前三块'),('输出校验','非空回答且引用编号存在'),('故障','模型异常明确返回，不伪造答案')],[4,10.5])
        elif section == 3:
            b.code(JAVA+'service/RagService.java','    public static List<String> split',end='    public List<RagSource> retrieve')
            b.code(JAVA+'service/AskService.java','        List<RagSource> sources = rag.retrieve','        String answer = model.ask')
            b.para('System提示明确禁止资料里的指令覆写系统规则，并要求资料未提供截止时间时说明待教师通知。User提示由当前课程名、问题及带编号的参考原文组成；模型返回后调用verifyCitations检查引用。')
            if (ROOT/'model-service/sdk_adapter.py').exists():
                b.code('model-service/sdk_adapter.py','def ask(question:',count=5)
                b.code('model-service/sdk_adapter.py','        status = getattr',count=4)
                b.para('SDK入口ask(question)实际调用官方Thread.run；只接受用户消息、推理和回答事件，随后检查完成状态并提取final_response。命名权限线程、文件读写探针和环境清理在connect中执行，未通过时禁用模型调用。')
        elif section == 4:
            b.para('Java构建日志e5-ask-build.log记录40条测试通过，包含RagServiceTest的6条和AskServiceTest、ModelClientTest合计11条。课程过滤故意移除后的测试失败保存在e5-course-filter-guard-failure.log，恢复后重新通过。单元测试覆盖检索及接口逻辑，不作为真实模型成功证据。')
            prompts = load_json('prompt-check.json')
            if prompts:
                b.para('System与User三次真实练习记录于prompt-check.json。前两次使用相同约束，分别不给原文和提供公开演示原文；第三次明确无资料不能生成编号引用，观察提示修正后的回答。')
                b.table(['练习','User内容','实际回答观察'],[
                    ('无上下文','只问报告何时交','回答无法确定日期；拒答带[1]，但输入未提供编号原文，引用不严谨。'),
                    ('提供上下文','[1] 五实验DOCX、源码SQL环境ZIP、截止待通知','正确概括两项交付，注明截止待教师通知，并引用[1]。'),
                    ('修正引用规则','仍不给原文；System要求无来源不编号','明确无法确定截止时间，没有生成不存在的[1]。')],[3,4.8,6.7])
            model = load_json('model-check.json')
            status = model.get('current_run_passed', model.get('passed', False))
            if model:
                b.para('真实模型及问答接口证据：model-check.json。当前记录的总体结果为' + ('通过' if status else '尚未通过') + '。详细结果以证据中的逐项观察为准，不将SDK单元测试或固定响应视为真实模型生成。')
                observations = []
                for test in model.get('tests', []):
                    details = test.get('details',{})
                    answer = details.get('answer','')
                    summary = re.sub(r'\*\*','',answer).replace('\n',' ')
                    if test.get('test') == 'MODEL-01':
                        summary = '回答五实验合并DOCX、源码与SQL配置和环境说明压缩包、本人推送后填写GitHub；注明“截止时间：尚未确认，待教师通知。[1]”。'
                    if test.get('test') == 'MODEL-UNAVAILABLE-01':
                        summary = '关闭模型桥后的实际HTTP状态：'+str(details.get('http',details.get('status','见原始证据')))+'；明确错误，未用固定答案替代。'
                    observations.append((test.get('test','检查'),test.get('result','未验证'),summary))
                b.table(['检查','结果','实际回答观察'],observations,[4.2,1.3,9])
                real = next((t for t in model.get('tests',[]) if t.get('test')=='MODEL-01'),{})
                if real:
                    sources = real.get('details',{}).get('sources',[])
                    for src in sources:
                        b.para('回答依据：'+str(src.get('title',''))+'；来源：'+str(src.get('source',''))+'；课程编号'+str(src.get('courseId',''))+'，片段编号'+str(src.get('chunkIndex',''))+'，相似度'+format(src.get('score',0),'.4f')+'。模型说明截止时间尚未确认，待教师通知，未补造日期。')
                fault = load_json('model-fault-check.json')
                if fault: b.para('模型故障实际检查记录于model-fault-check.json；关闭模型桥后端返回明确错误，未返回替代答案。')
                unavailable = next((t for t in model.get('tests',[]) if t.get('test')=='MODEL-UNAVAILABLE-01'),{})
                if unavailable:
                    b.para('实际关闭本地8787模型桥后重新请求，MODEL-UNAVAILABLE-01记录明确错误，未返回固定替代答案；随后恢复模型桥。此结果与无资料拒答区分。')
            else:
                b.para('真实模型调用目前尚未完成验收：本地Codex SDK桥的安全沙箱仍在修复，尚无model-check.json或rag.png证明端到端问答成功。实验五AC-E5-01与AC-E5-02保持待验证；本报告不展示合成成功截图或虚构幻觉案例。')
            if (EVIDENCE/'rag.png').exists():
                b.image('rag.png','图5 1 当前课程问题、真实模型回答与引用（真实截图局部）', (304,714,1064,1244))
                b.image('rag.png','图5 2 同一次回答展开的原文来源（同一真实截图局部）', (325,1250,1368,1562))
        else:
            reflection = EVIDENCE/'ai-reflection.md'
            if reflection.exists():
                b.para('真实提示练习暴露了无原文时拒答仍带[1]的问题；改为“只引用实际提供的编号，无资料不编号”后第三次回答不再虚构引用。正式RAG先检索、无资料提前拒答，并返回原文供核对；编号存在仍不能证明每项事实都受原文支持。三次练习与返回见prompt-check.json，详细反思见ai-reflection.md。')
                b.para('工程中Windows权限探针启动失败、正常userMessage被白名单误拒绝，分别通过专用空目录及必要系统环境、放行正常消息解决。探针只证明本机公开演示的有限隔离。故障测试首次问题过短而走无资料拒答；独立作者改用已命中的问题，真实停桥后得到503，保留原失败并恢复服务。')
                b.para('真实生成使用官方Python SDK及ChatGPT登录；兼容API Key模式未实际验证。本人应对照原文逐项核对回答，并能解释页面到数据库、JWT授权、检索到模型的流程；报告未声称本人已逐行理解。截止时间保留待教师通知，GitHub由本人推送后填写。')
            else:
                model = load_json('model-check.json')
                if model.get('current_run_passed',model.get('passed',False)):
                    b.para('AI协助实现检索、提示拼接和SDK桥。真实模型回答与本次演示原文吻合，保留[1]来源并说明截止待教师通知；本次未观察到可据以记录的模型幻觉，不编造幻觉案例。')
                    b.para('实际故障包括初期Windows沙箱探针启动失败，以及事件白名单误拒绝正常userMessage回显。采用专用空目录、子进程运行目录与必要系统环境后，读写边界探针通过；将正常消息类型纳入白名单后完成真实调用。该边界是本机公开演示的有限隔离，不代表整机文件读取完全隔离。')
                    if load_json('prompt-check.json'):
                        b.para('独立提示练习中，无原文时拒答仍带[1]，说明只要求引用不足以保证来源有效。正式RAG流程在无检索原文时提前拒答，不调用模型；有原文时返回片段供逐项核对。学习重点是先查资料，再回答，并如实保留未知截止时间。')
                else:
                    b.para('AI协助实现检索和提示拼接，测试检查课程过滤、无资料和引用规则。真实模型验收尚未完成，不能据此编造模型幻觉或声称问答正确。')
                b.para('引用编号存在只能证明指向本次来源，不能保证每项事实得到支持。后续核对须将答案逐项对应原文；无资料、跨课程和泄密问题应拒绝，模型故障应显示明确错误。')
            b.para('AI生成、AI检查和实际运行证据在本报告中分别记录。报告和源码用于学习与提交准备；本人还需阅读并复述关键流程。本地GitHub地址待自行推送后填写，未产生远程PR或云端CI通过结果。')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--pages',type=Path)
    args = parser.parse_args()
    if digest(REFERENCE) != '8b88124b5d66f0039b9d19deb07a22ca463b54fd284834de3a7fcafd18dd5483':
        raise SystemExit('Reference hash does not match artifact contract')
    make(json.loads(args.pages.read_text(encoding='utf-8-sig')) if args.pages else {})
